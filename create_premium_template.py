"""
create_premium_template.py
==========================
Creates the Premium Corporate MRC Business Review PowerPoint template.

This template uses a sophisticated "SaaS Dashboard" color palette:
- Deep Sapphire Blue (#1E3A8A) - dark text
- Sky Blue (#0EA5E9) - accents
- Forest Green (#10B981) - success/positive
- Amber Gold (#F59E0B) - warnings
- Deep Rose (#E11D48) - alerts
- Soft Grey-Blue (#F4F6F9) - backgrounds

Usage:
    python create_premium_template.py

Output:
    template/MRC_Business_Review_Template_Premium.pptx
"""

import os
import shutil
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def create_premium_template(
    output_dir: str = "template",
    filename: str = "MRC_Business_Review_Template_Premium.pptx",
    backup: bool = True
) -> str:
    """
    Create MRC Business Review Premium Corporate template PPTX.
    
    Args:
        output_dir: Directory to save template
        filename: Template filename
        backup: Backup existing template before creating new
    
    Returns:
        Path to created template file
    """
    
    # Ensure output directory exists
    os.makedirs(output_dir, exist_ok=True)
    template_path = os.path.join(output_dir, filename)
    
    # Backup existing template
    if os.path.exists(template_path) and backup:
        backup_dir = os.path.join(output_dir, "backups")
        os.makedirs(backup_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = os.path.join(
            backup_dir,
            f"MRC_Business_Review_Template_Premium_backup_{timestamp}.pptx"
        )
        shutil.copy2(template_path, backup_path)
        print(f"✓ Backed up existing template: {backup_path}")
    
    # Create presentation
    prs = Presentation()
    
    # Set slide dimensions (16:9 widescreen)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.50)
    
    # Define colors (Premium Corporate Palette)
    SKY_BLUE = RGBColor(14, 165, 233)        # #0EA5E9 - Primary accent
    DEEP_SAPPHIRE = RGBColor(30, 58, 138)    # #1E3A8A - Dark text
    WHITE = RGBColor(255, 255, 255)
    SOFT_GREY_BLUE = RGBColor(244, 246, 249) # #F4F6F9 - Background
    GREY_LIGHT = RGBColor(226, 232, 240)     # #E2E8F0 - Dividers
    GREY_MID = RGBColor(100, 116, 139)       # #64748B - Mid tone
    
    print("\n" + "="*60)
    print("CREATING MRC PREMIUM CORPORATE TEMPLATE")
    print("="*60)
    print("Theme: SaaS Dashboard with muted corporate colors")
    print(f"Primary: Sky Blue (#0EA5E9)")
    print(f"Dark: Deep Sapphire (#1E3A8A)")
    print(f"Background: Soft Grey-Blue (#F4F6F9)")
    
    # ──────────────────────────────────────────────────────────
    # Slide 1: Title Slide Template
    # ──────────────────────────────────────────────────────────
    print("\n[1/3] Creating title slide template...")
    
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    slide1.background.fill.solid()
    slide1.background.fill.fore_color.rgb = WHITE
    
    # Left Deep Sapphire sidebar
    sidebar = slide1.shapes.add_shape(
        1,  # Rectangle
        Inches(0),
        Inches(0),
        Inches(4.5),
        Inches(7.50)
    )
    sidebar.fill.solid()
    sidebar.fill.fore_color.rgb = DEEP_SAPPHIRE
    sidebar.line.fill.background()
    
    # Title area (template text)
    tb_title = slide1.shapes.add_textbox(
        Inches(0.35),
        Inches(2.0),
        Inches(3.8),
        Inches(3.2)
    )
    tf = tb_title.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "MRC TÜRKİYE"
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = WHITE
    
    p1 = tf.add_paragraph()
    p1.text = "[YEAR]"
    p1.font.bold = True
    p1.font.size = Pt(52)
    p1.font.color.rgb = WHITE
    
    p2 = tf.add_paragraph()
    p2.text = ""
    p2.font.size = Pt(10)
    
    p3 = tf.add_paragraph()
    p3.text = "[MONTH] Financial Results"
    p3.font.size = Pt(18)
    p3.font.color.rgb = WHITE
    
    # Subtitle area
    tb_subtitle = slide1.shapes.add_textbox(
        Inches(5.0),
        Inches(3.0),
        Inches(7.8),
        Inches(1.5)
    )
    p_s = tb_subtitle.text_frame.paragraphs[0]
    p_s.text = "Monthly Business Review · [MONTH] [YEAR]"
    p_s.font.size = Pt(14)
    p_s.font.color.rgb = GREY_MID
    
    # Premium tag
    tb_tag = slide1.shapes.add_textbox(
        Inches(5.0),
        Inches(0.3),
        Inches(7.8),
        Inches(0.6)
    )
    p_tag = tb_tag.text_frame.paragraphs[0]
    p_tag.text = "Premium Corporate Edition"
    p_tag.font.size = Pt(10)
    p_tag.font.italic = True
    p_tag.font.color.rgb = GREY_MID
    
    print("   ✓ Title slide template created")
    
    # ──────────────────────────────────────────────────────────
    # Slide 2: Content Slide Template
    # ──────────────────────────────────────────────────────────
    print("[2/3] Creating content slide template...")
    
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = SOFT_GREY_BLUE
    
    # White content area
    content_bg = slide2.shapes.add_shape(
        1,
        Inches(0),
        Inches(0.9),
        Inches(13.33),
        Inches(6.6)
    )
    content_bg.fill.solid()
    content_bg.fill.fore_color.rgb = WHITE
    content_bg.line.fill.background()
    
    # Header
    tb_header = slide2.shapes.add_textbox(
        Inches(0.45),
        Inches(0.18),
        Inches(12.43),
        Inches(0.72)
    )
    tf_h = tb_header.text_frame
    tf_h.word_wrap = False
    p_h = tf_h.paragraphs[0]
    p_h.text = "[CHART TITLE]"
    p_h.font.bold = True
    p_h.font.size = Pt(28)
    p_h.font.color.rgb = DEEP_SAPPHIRE
    
    # Header divider line (thicker and more prominent)
    div_line = slide2.shapes.add_shape(
        1,
        Inches(0.45),
        Inches(0.95),
        Inches(12.43),
        Inches(0.03)
    )
    div_line.fill.solid()
    div_line.fill.fore_color.rgb = SKY_BLUE
    div_line.line.fill.background()
    
    # Subtitle
    tb_subtitle2 = slide2.shapes.add_textbox(
        Inches(0.45),
        Inches(1.02),
        Inches(12.43),
        Inches(0.30)
    )
    p_sub = tb_subtitle2.text_frame.paragraphs[0]
    p_sub.text = "[SUBTITLE]"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = GREY_MID
    
    # Chart placeholder (light soft background)
    chart_placeholder = slide2.shapes.add_shape(
        1,
        Inches(0.45),
        Inches(1.32),
        Inches(8.0),
        Inches(5.0)
    )
    chart_placeholder.fill.solid()
    chart_placeholder.fill.fore_color.rgb = SOFT_GREY_BLUE
    chart_placeholder.line.color.rgb = GREY_LIGHT
    chart_placeholder.line.width = Pt(1)
    
    # Chart placeholder text
    tb_chart = slide2.shapes.add_textbox(
        Inches(0.45),
        Inches(3.5),
        Inches(8.0),
        Inches(0.5)
    )
    p_chart = tb_chart.text_frame.paragraphs[0]
    p_chart.text = "[CHART IMAGE]"
    p_chart.font.size = Pt(12)
    p_chart.font.color.rgb = GREY_MID
    p_chart.alignment = PP_ALIGN.CENTER
    
    # Callout box template (premium style)
    callout = slide2.shapes.add_shape(
        5,  # Rounded rectangle
        Inches(8.65),
        Inches(1.32),
        Inches(4.23),
        Inches(5.0)
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = SOFT_GREY_BLUE
    callout.line.color.rgb = SKY_BLUE
    callout.line.width = Pt(1.5)
    
    tb_callout = slide2.shapes.add_textbox(
        Inches(8.85),
        Inches(1.52),
        Inches(3.83),
        Inches(4.6)
    )
    p_callout = tb_callout.text_frame.paragraphs[0]
    p_callout.text = "[KEY METRICS]"
    p_callout.font.size = Pt(11)
    p_callout.font.color.rgb = GREY_MID
    
    print("   ✓ Content slide template created")
    
    # ──────────────────────────────────────────────────────────
    # Slide 3: Thank You Slide Template
    # ──────────────────────────────────────────────────────────
    print("[3/3] Creating thank you slide template...")
    
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    slide3.background.fill.solid()
    slide3.background.fill.fore_color.rgb = SOFT_GREY_BLUE
    
    # Top accent bar (Sky Blue)
    top_bar = slide3.shapes.add_shape(
        1,
        Inches(0),
        Inches(0),
        Inches(13.33),
        Inches(0.2)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = SKY_BLUE
    top_bar.line.fill.background()
    
    # Bottom accent bar (Deep Sapphire)
    bottom_bar = slide3.shapes.add_shape(
        1,
        Inches(0),
        Inches(7.30),
        Inches(13.33),
        Inches(0.2)
    )
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = DEEP_SAPPHIRE
    bottom_bar.line.fill.background()
    
    # White content area in middle
    content_area = slide3.shapes.add_shape(
        1,
        Inches(0),
        Inches(0.2),
        Inches(13.33),
        Inches(7.1)
    )
    content_area.fill.solid()
    content_area.fill.fore_color.rgb = WHITE
    content_area.line.fill.background()
    
    # Main text (Deep Sapphire for premium feel)
    tb_main = slide3.shapes.add_textbox(
        Inches(1),
        Inches(2.2),
        Inches(11.33),
        Inches(2)
    )
    p_main = tb_main.text_frame.paragraphs[0]
    p_main.text = "THANK YOU"
    p_main.font.bold = True
    p_main.font.size = Pt(80)
    p_main.font.color.rgb = DEEP_SAPPHIRE
    p_main.alignment = PP_ALIGN.CENTER
    
    # Footer text
    tb_footer = slide3.shapes.add_textbox(
        Inches(1),
        Inches(4.8),
        Inches(11.33),
        Inches(0.8)
    )
    p_footer = tb_footer.text_frame.paragraphs[0]
    p_footer.text = "MRC TÜRKİYE  ·  [MONTH] [YEAR]  ·  www.mrc-international.com"
    p_footer.font.size = Pt(11)
    p_footer.font.color.rgb = GREY_MID
    p_footer.alignment = PP_ALIGN.CENTER
    
    print("   ✓ Thank you slide template created")
    
    # ──────────────────────────────────────────────────────────
    # Save template
    # ──────────────────────────────────────────────────────────
    print("\n" + "="*60)
    print("SAVING PREMIUM TEMPLATE")
    print("="*60)
    
    try:
        prs.save(template_path)
        file_size = os.path.getsize(template_path)
        print(f"\n✓ Premium template saved: {template_path}")
        print(f"✓ File size: {file_size:,} bytes ({file_size/1024:.1f} KB)")
        print(f"✓ Slides: {len(prs.slides)}")
        print(f"✓ Dimensions: 13.33\" × 7.50\" (16:9 widescreen)")
        print(f"✓ Theme: Premium Corporate (SaaS Dashboard)")
        print(f"✓ Primary Colors: Sky Blue, Deep Sapphire, Soft Grey-Blue")
        
        print("\n" + "="*60)
        print("PREMIUM TEMPLATE READY")
        print("="*60)
        print(f"\nTo use this template in a script:")
        print(f"  from pptx import Presentation")
        print(f"  prs = Presentation('{template_path}')")
        print(f"\nTo switch to premium design config:")
        print(f"  config.get('design_premium.COLORS.Primary.Hex')")
        print(f"  → '#0EA5E9' (Sky Blue)")
        
        return template_path
        
    except Exception as e:
        print(f"\n✗ ERROR saving template: {e}")
        return None


if __name__ == '__main__':
    create_premium_template()
